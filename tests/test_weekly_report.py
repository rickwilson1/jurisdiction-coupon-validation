"""Weekly coupon activity report: week bounds, aggregation, rendering, xlsx.

Runs with `python tests/test_weekly_report.py` (needs openpyxl); pytest will
also collect it. Imports `dispatch.weekly_report` directly so no Firestore,
Graph, or reportlab dependencies are needed.
"""

import os
import sys
from datetime import UTC, date, datetime
from io import BytesIO

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from dispatch import weekly_report as wr  # noqa: E402

_FAILURES = []
_RUN = []


def check(label, actual, expected):
    _RUN.append(label)
    if actual == expected:
        print(f"  PASS  {label}")
    else:
        print(f"  FAIL  {label}: got {actual!r}, expected {expected!r}")
        _FAILURES.append(label)


def doc(number, order_date, code, routing, qty, material="Compost A - Compost", **extra):
    d = {
        "order_number": number,
        "order_date": order_date,
        "coupon_code": code,
        "routing": routing,
        "region": "oc",
        "shipping_method": "Frank R. Bowerman Landfill Pick Up",
        "material": material,
        "total_qty": qty,
        "customer_name": "Test Customer",
        "customer_email": "test@example.com",
        "customer_phone": "555-0100",
        "shipping_address": "1 Main St, Irvine, CA 92618",
        "status": "success",
        "processed_at": datetime(2026, 9, 1, 12, 0, tzinfo=UTC),
    }
    d.update(extra)
    return d


DOCS = [
    doc("A1", "8/1/2026 9:00:00 AM PT", "CITYSAACOM26", "pickup_staff_load", 8.0),
    doc("A100", "8/26/2026 9:00:00 AM PT", "CITYSAACOM26", "pickup_self_load", 0.5),
    doc(
        "A101",
        "8/27/2026 9:00:00 AM PT",
        "CITYANACM26",
        "pickup_staff_load",
        8.0,
        "Compost B - Cover Mulch",
    ),
    doc("A102", "8/31/2026 9:00:00 AM PT", "CITYBUPCOM26", "pickup_self_load", 0.5),
    doc("A103", "9/2/2026 9:00:00 AM PT", "CITYSCLCOM26", "pickup_staff_load", 5.0),
    doc("A104", "9/6/2026 9:00:00 AM PT", "CITYIRVCOM26", "pickup_self_load", 1.0),
    doc("A105", "9/13/2026 9:00:00 AM PT", "CITYNPBCOM26", "pickup_self_load", 2.0),
    doc(
        "A106",
        "5/4/2026 9:00:00 AM PT",
        "IRVINECOM26",
        "delivery",
        3.0,
        "Compost",
        shipping_method="Frank R. Bowerman / Bee Canyon Pick Up Delivery",
    ),
    # No order_date; falls back to processed_at (Sep 1 05:00 UTC = Aug 31 PT).
    doc(
        "A107",
        None,
        "CITYPLACOM26",
        "pickup_self_load",
        0.5,
        processed_at=datetime(2026, 9, 1, 5, 0, tzinfo=UTC),
    ),
    # Future-dated relative to the report week; must be excluded from to-date.
    doc("A108", "9/15/2026 9:00:00 AM PT", "CITYSAACOM26", "pickup_self_load", 1.0),
]


print("\nreport_week")
check(
    "Monday gives the week that ended yesterday",
    wr.report_week(date(2026, 9, 14)),
    (date(2026, 9, 7), date(2026, 9, 13)),
)
check(
    "Sunday gives the previous complete week",
    wr.report_week(date(2026, 9, 13)),
    (date(2026, 8, 31), date(2026, 9, 6)),
)
check(
    "Wednesday gives the prior Mon-Sun",
    wr.report_week(date(2026, 9, 16)),
    (date(2026, 9, 7), date(2026, 9, 13)),
)

print("\nparse_order_date")
check("CIMcloud format", wr.parse_order_date("4/16/2026 11:38:23 AM PT"), date(2026, 4, 16))
check("garbage returns None", wr.parse_order_date("soon"), None)
check("None returns None", wr.parse_order_date(None), None)

print("\njurisdiction_for")
check(
    "master wins",
    wr.jurisdiction_for("CITYSAACOM26", {"CITYSAACOM26": "City of Santa Ana"}),
    "City of Santa Ana",
)
check("fallback abbreviation", wr.jurisdiction_for("CITYSAACM26"), "Santa Ana")
check("legacy Irvine code", wr.jurisdiction_for("IRVINECOM26"), "Irvine")
check("Sacramento bag code", wr.jurisdiction_for("CITYSACCOMB26"), "Sacramento")
check("unknown returns the code", wr.jurisdiction_for("CITYZZZCOM26"), "CITYZZZCOM26")

VENTURA = doc("A200", "9/9/2026 9:00:00 AM PT", "CITYVCOM26", "delivery", 5.0, region="ventura")

print("\nnormalize_events")
orders = wr.normalize_events(DOCS + [VENTURA], None, excluded={"A1"})
check("test order excluded", "A1" in {o.order_number for o in orders}, False)
check("default region filter drops Ventura", "A200" in {o.order_number for o in orders}, False)
check("count", len(orders), 9)
check(
    "region=None keeps everything",
    len(wr.normalize_events(DOCS + [VENTURA], None, excluded={"A1"}, region=None)),
    10,
)
check(
    "explicit region match is case-insensitive",
    [
        o.order_number
        for o in wr.normalize_events([VENTURA], None, excluded=set(), region="VENTURA")
    ],
    ["A200"],
)
check(
    "processed_at fallback lands on Pacific date",
    next(o for o in orders if o.order_number == "A107").order_date,
    date(2026, 8, 31),
)
check("sorted by date", [o.order_number for o in orders][:2], ["A106", "A100"])

print("\nbuild_report")
r = wr.build_report(
    orders, date(2026, 9, 7), date(2026, 9, 13), generated_at=datetime(2026, 9, 14, 7, 0)
)
check("program start defaults to launch", r.program_start, date(2026, 8, 28))
check("week orders", r.week.orders, 1)
check("week CY", r.week.cubic_yards, 2.0)
check("week jurisdictions", r.week.jurisdictions, ["Newport Beach"])
check("prior week orders", r.prior_week.orders, 4)
check("prior week CY", r.prior_week.cubic_yards, 7.0)
check(
    "pre-launch bucket (A106, A100, A101)",
    (r.pre_launch.orders, r.pre_launch.cubic_yards),
    (3, 11.5),
)
check("pre-launch label", r.pre_launch.label, "Pre-launch (before Aug 28)")
check("to-date excludes pre-launch and future", r.to_date.orders, 5)
check("to-date CY", r.to_date.cubic_yards, 9.0)
check("to-date routing", (r.to_date.self_load, r.to_date.staff_load, r.to_date.delivery), (4, 1, 0))
check(
    "weekly table starts at launch week",
    [b.label for b in r.weekly],
    ["Aug 24 to 30", "Aug 31 to Sep 6", "Sep 7 to 13"],
)
check("launch week shows only post-launch orders", r.weekly[0].orders, 0)
check(
    "jurisdiction order (orders desc, CY desc, name)",
    [b.label for b in r.by_jurisdiction],
    ["San Clemente", "Newport Beach", "Irvine", "Buena Park", "Placentia"],
)
check("pre-launch jurisdictions absent", "Anaheim" in [b.label for b in r.by_jurisdiction], False)
check(
    "monthly spans launch month to report month",
    [b.label for b in r.monthly],
    ["August 2026", "September 2026"],
)
check("August counts launch-onward only", (r.monthly[0].orders, r.monthly[0].cubic_yards), (2, 1.0))
check("material groups", [(b.label, b.cubic_yards) for b in r.by_material], [("Compost", 9.0)])
check("financials flagged unavailable", r.financials_available, False)
check("first order date", r.first_order_date, date(2026, 8, 31))

print("\nbuild_report with explicit program_start")
r_all = wr.build_report(
    orders,
    date(2026, 9, 7),
    date(2026, 9, 13),
    generated_at=datetime(2026, 9, 14, 7, 0),
    program_start=date(2026, 1, 1),
)
check("no pre-launch when start precedes data", r_all.pre_launch.orders, 0)
check("to-date includes everything", r_all.to_date.orders, 8)
check("eight weekly buckets when launch is old", len(r_all.weekly), 8)
check(
    "Irvine groups legacy and new code",
    next(b for b in r_all.by_jurisdiction if b.label == "Irvine").codes,
    ["IRVINECOM26", "CITYIRVCOM26"],
)

print("\nrendering")
subj = wr.subject_line(r)
check(
    "subject",
    subj,
    "OCWR Coupon Activity, Week of September 7 to 13, 2026: 1 order, 2 CY; 5 orders / 9 CY to date",
)
html_body = wr.render_html(r)
check(
    "html has no PII",
    ("test@example.com" in html_body)
    or ("Test Customer" in html_body)
    or ("1 Main St" in html_body),
    False,
)
check("html has week table", "This week" in html_body, True)
check("html has jurisdiction total row", ">Total<" in html_body, True)
check("html titled OCWR", "OCWR Coupon Program Activity" in html_body, True)
check("html has pre-launch row", "Pre-launch (before Aug 28)" in html_body, True)
check("html states launch date", "launched August 28, 2026" in html_body, True)
check("html has no pre-launch row when none", "Pre-launch" in wr.render_html(r_all), False)
check("html has no notes section", "Notes</p>" in html_body, False)
check("html uses inline font stack", "font-family:Aptos" in html_body, True)
check("html has no external css", "<link" in html_body or "<style" in html_body, False)
text = wr.render_text(r)
check("text has summary", "1 coupon order for 2 cubic yards" in text, True)
check("text has no PII", "test@example.com" in text, False)
summary = wr._summary_sentences(r)
check("delta sentence", summary[1], "That is down from 4 orders and 7 CY in the prior week.")

print("\nempty week")
r0 = wr.build_report(
    orders, date(2026, 8, 3), date(2026, 8, 9), generated_at=datetime(2026, 8, 10, 7, 0)
)
check(
    "zero-week summary",
    wr._summary_sentences(r0)[0],
    "No coupon orders were recorded for August 3 to 9, 2026.",
)
check("zero-week html renders", "No coupon orders" in wr.render_html(r0), True)
check("week before launch has empty to-date", r0.to_date.orders, 0)
check("week before launch still buckets pre-launch through week end", r0.pre_launch.orders, 1)
check("no monthly rows before launch", r0.monthly, [])
check("no weekly rows before launch", r0.weekly, [])

print("\nxlsx")
data = wr.build_xlsx(r)
wb = load_workbook(BytesIO(data))
check(
    "sheets",
    wb.sheetnames,
    ["Week orders", "All orders", "Weekly", "By jurisdiction", "Monthly", "Material and site"],
)
ws = wb["All orders"]
check("all orders rows (launch onward)", ws.max_row - 1, 5)
check("all orders carries customer email column", "Email" in [c.value for c in ws[1]], True)
check("header font Aptos", ws["A1"].font.name, "Aptos")
check("body font Aptos 11", (ws["A2"].font.name, ws["A2"].font.size), ("Aptos", 11.0))
check("order date is a date cell", isinstance(ws["B2"].value, (date, datetime)), True)
check("week orders rows", wb["Week orders"].max_row - 1, 1)
check(
    "attachment filename",
    wr.attachment_filename(r),
    "Coupon_Activity_2026-09-07_to_2026-09-13.xlsx",
)

print("\nrecipients")
os.environ.pop("WEEKLY_REPORT_TO", None)
check("default recipient count", len(wr.recipients_from_env()), 13)
os.environ["WEEKLY_REPORT_TO"] = "a@agromin.com, b@agromin.com,"
check("env override", wr.recipients_from_env(), ["a@agromin.com", "b@agromin.com"])
os.environ.pop("WEEKLY_REPORT_TO", None)
os.environ["WEEKLY_REPORT_PROGRAM_START"] = "2026-09-01"
check("program start env override", wr.program_start_from_env(), date(2026, 9, 1))
os.environ.pop("WEEKLY_REPORT_PROGRAM_START", None)
check("program start default", wr.program_start_from_env(), date(2026, 8, 28))
os.environ.pop("WEEKLY_REPORT_REGION", None)
check("region default", wr.region_from_env(), "oc")
os.environ["WEEKLY_REPORT_REGION"] = ""
check("empty region env means no filter", wr.region_from_env(), None)
os.environ.pop("WEEKLY_REPORT_REGION", None)
os.environ["WEEKLY_REPORT_EXCLUDE"] = "a1,T9"
check("exclude env uppercases", wr.excluded_orders_from_env(), {"A1", "T9"})
os.environ.pop("WEEKLY_REPORT_EXCLUDE", None)

print(f"\n{len(_RUN) - len(_FAILURES)}/{len(_RUN)} passed")
if _FAILURES:
    sys.exit(1)
